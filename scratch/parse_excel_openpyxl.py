import openpyxl

def main():
    path = "scratch/threat_intel_input.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    print("Sheets in workbook:", wb.sheetnames)
    
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"\n--- Sheet: {name} ---")
        print(f"Max row: {sheet.max_row}, Max column: {sheet.max_column}")
        
        # Read header row
        headers = [cell.value for cell in sheet[1]]
        print("Headers:", headers)
        
        # Print first few rows
        print("Data rows:")
        count = 0
        for r in range(2, sheet.max_row + 1):
            row_vals = [cell.value for cell in sheet[r]]
            if any(row_vals):
                print(f"  Row {r}: {row_vals}")
                count += 1
                if count >= 10:
                    break

if __name__ == "__main__":
    main()
