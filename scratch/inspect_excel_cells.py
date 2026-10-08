import openpyxl

def main():
    path = "scratch/threat_intel_input.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"\n--- Sheet: {name} ---")
        # Iterate over all rows and cols to see if there are any cells with values
        cells_with_values = []
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(row=r, column=c).value
                if val is not None:
                    cells_with_values.append((r, c, val))
        
        print(f"Total cells with values: {len(cells_with_values)}")
        for r, c, val in cells_with_values[:20]:
            print(f"  Cell({r},{c}): {val}")

if __name__ == "__main__":
    main()
