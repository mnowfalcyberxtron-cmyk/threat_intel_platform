import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ── Parse the raw TSV data ────────────────────────────────────────────────────
rows = []
with open("raw_cve_data.txt", encoding="utf-8") as f:
    lines = f.readlines()

headers = [h.strip() for h in lines[0].split("\t")]
for line in lines[1:]:
    line = line.strip()
    if not line:
        continue
    parts = line.split("\t")
    # Pad short rows
    while len(parts) < len(headers):
        parts.append("")
    row = {headers[i]: parts[i].strip() for i in range(len(headers))}
    rows.append(row)

# ── Column order as user wants ────────────────────────────────────────────────
COL_ORDER = [
    "CVE ID",
    "TITLE",
    "XTRON SCORE",
    "IMPACT",
    "PATCH AVAILABILITY",
    "Score",
    "Severity",
    "VENDOR",
    "PRODUCT",
    "AFFECTED VERSION",
    "FIXED VERSION",
]

# ── Severity → color fill ─────────────────────────────────────────────────────
def sev_fill(sev):
    s = str(sev).strip().lower()
    if s == "critical":
        return PatternFill("solid", fgColor="C00000")   # dark red
    if s == "high":
        return PatternFill("solid", fgColor="FF4444")   # red
    if s == "medium":
        return PatternFill("solid", fgColor="FFA500")   # orange
    if s == "low":
        return PatternFill("solid", fgColor="FFFF00")   # yellow
    return PatternFill("solid", fgColor="D9D9D9")        # grey N/A

def patch_fill(val):
    v = str(val).strip().upper()
    if v == "TRUE":
        return PatternFill("solid", fgColor="70AD47")   # green
    if v == "FALSE":
        return PatternFill("solid", fgColor="FF4444")   # red
    return None

# ── Score → severity label (if missing) ──────────────────────────────────────
def score_to_sev(score_str):
    try:
        s = float(score_str)
        if s >= 9.0: return "Critical"
        if s >= 7.0: return "High"
        if s >= 4.0: return "Medium"
        return "Low"
    except:
        return "N/A"

# ── Build workbook ────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ICS CVE Advisory"

# Header style
header_fill  = PatternFill("solid", fgColor="1F3864")   # navy
header_font  = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
thin         = Side(style="thin", color="BFBFBF")
border       = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.row_dimensions[1].height = 32

# Write headers
for ci, col in enumerate(COL_ORDER, start=1):
    cell = ws.cell(row=1, column=ci, value=col)
    cell.fill      = header_fill
    cell.font      = header_font
    cell.alignment = center_align
    cell.border    = border

# ── Column widths ─────────────────────────────────────────────────────────────
col_widths = {
    "CVE ID":             18,
    "TITLE":              45,
    "XTRON SCORE":        13,
    "IMPACT":             55,
    "PATCH AVAILABILITY": 18,
    "Score":              10,
    "Severity":           12,
    "VENDOR":             22,
    "PRODUCT":            32,
    "AFFECTED VERSION":   35,
    "FIXED VERSION":      40,
}
for ci, col in enumerate(COL_ORDER, start=1):
    ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(col, 20)

# ── Write data rows ───────────────────────────────────────────────────────────
alt_fill = PatternFill("solid", fgColor="EEF3F7")   # light blue-grey for alternating

for ri, row in enumerate(rows, start=2):
    ws.row_dimensions[ri].height = 52

    sev   = row.get("Severity", "").strip()
    patch = row.get("PATCH AVAILABILITY", "").strip()
    score = row.get("Score", "").strip()

    # If severity blank, derive from score
    if not sev or sev.upper() == "N/A":
        sev = score_to_sev(score)

    is_alt = (ri % 2 == 0)
    row_base_fill = alt_fill if is_alt else PatternFill("solid", fgColor="FFFFFF")

    for ci, col in enumerate(COL_ORDER, start=1):
        val = row.get(col, "")

        # Normalise severity value
        if col == "Severity":
            val = sev

        cell = ws.cell(row=ri, column=ci, value=val)
        cell.border    = border

        # Alignment
        if col in ("CVE ID", "XTRON SCORE", "Score", "Severity", "PATCH AVAILABILITY"):
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = left_align

        # Font
        cell.font = Font(name="Calibri", size=10)

        # Cell fills
        if col == "Severity":
            cell.fill = sev_fill(sev)
            cell.font = Font(name="Calibri", size=10, bold=True,
                             color="FFFFFF" if sev.lower() in ("critical","high") else "000000")
        elif col == "PATCH AVAILABILITY":
            pf = patch_fill(patch)
            cell.fill = pf if pf else row_base_fill
            cell.font = Font(name="Calibri", size=10, bold=True,
                             color="FFFFFF" if patch.upper() == "TRUE" else
                             ("FFFFFF" if patch.upper() == "FALSE" else "000000"))
        elif col == "XTRON SCORE":
            try:
                xs = int(val)
                if xs >= 70:
                    cell.fill = PatternFill("solid", fgColor="C00000")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
                elif xs >= 50:
                    cell.fill = PatternFill("solid", fgColor="FFA500")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="000000")
                else:
                    cell.fill = PatternFill("solid", fgColor="FFFF00")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="000000")
            except:
                cell.fill = row_base_fill
        else:
            cell.fill = row_base_fill

# ── Freeze top row, enable auto-filter ───────────────────────────────────────
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COL_ORDER))}1"

# ── Save ─────────────────────────────────────────────────────────────────────
out_path = "ICS_CVE_Advisory_Final.xlsx"
wb.save(out_path)
print(f"Saved {out_path}  ({len(rows)} rows)")
