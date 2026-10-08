#!/usr/bin/env python3
"""Diagnostic script for ICS, Dark Web, and Breach Market issues."""
import asyncio
from database.db import Database
from config import settings

async def diagnose():
    db = Database()
    await db.initialize()
    
    print("=" * 80)
    print("DIAGNOSTIC REPORT - Data Collection Issues")
    print("=" * 80)
    print()
    
    # ===== ISSUE 1: ICS Advisories missing fields =====
    print("📋 ISSUE 1: ICS ADVISORIES - PATCH, POC, CVSS SEVERITY, VENDOR, APPLICATION")
    print("-" * 80)
    try:
        async with db._conn.execute(
            "SELECT COUNT(*) as cnt FROM ics_advisories"
        ) as cur:
            row = await cur.fetchone()
            ics_count = row["cnt"] if row else 0
        
        if ics_count == 0:
            print("❌ NO ICS ADVISORIES IN DATABASE!")
            print("   Status: ICS sync job hasn't run yet or failed")
            print("   Action: Wait for ics_sync job (runs every 60 min)")
            print("   Or manually trigger: POST /api/advisory/ics/sync")
        else:
            print(f"✅ {ics_count:,} ICS advisories in database")
            
            # Check if required fields are populated
            async with db._conn.execute(
                """SELECT 
                   COUNT(*) as total,
                   SUM(CASE WHEN patch_availability != 'Unknown' THEN 1 ELSE 0 END) as has_patch,
                   SUM(CASE WHEN poc_availability != 'Unknown' THEN 1 ELSE 0 END) as has_poc,
                   SUM(CASE WHEN severity != 'Unknown' AND severity != '' THEN 1 ELSE 0 END) as has_severity,
                   SUM(CASE WHEN vendor != 'Unknown' AND vendor != '' THEN 1 ELSE 0 END) as has_vendor,
                   SUM(CASE WHEN product != 'Unknown' AND product != '' THEN 1 ELSE 0 END) as has_product
                FROM ics_advisories"""
            ) as cur:
                stats = await cur.fetchone()
            
            total = stats["total"] or 0
            has_patch = stats["has_patch"] or 0
            has_poc = stats["has_poc"] or 0
            has_severity = stats["has_severity"] or 0
            has_vendor = stats["has_vendor"] or 0
            has_product = stats["has_product"] or 0
            
            print(f"   • Patch availability:      {has_patch}/{total} ({100*has_patch/total:.1f}%)")
            print(f"   • POC availability:        {has_poc}/{total} ({100*has_poc/total:.1f}%)")
            print(f"   • CVSS severity:           {has_severity}/{total} ({100*has_severity/total:.1f}%)")
            print(f"   • Affected vendor:         {has_vendor}/{total} ({100*has_vendor/total:.1f}%)")
            print(f"   • Affected application:    {has_product}/{total} ({100*has_product/total:.1f}%)")
            
            if has_severity < total * 0.5:
                print("   ⚠️  Less than 50% have CVSS severity - needs data enrichment")
    except Exception as e:
        print(f"❌ Error checking ICS data: {e}")
    
    print()
    
    # ===== ISSUE 2: Dark Web Manager not scanning =====
    print("🕷️  ISSUE 2: DARK WEB MANAGER - NOT SCANNING")
    print("-" * 80)
    try:
        async with db._conn.execute(
            "SELECT COUNT(*) as cnt FROM onion_sites"
        ) as cur:
            row = await cur.fetchone()
            onion_count = row["cnt"] if row else 0
        
        if onion_count == 0:
            print("❌ NO ONION SITES CONFIGURED!")
            print("   Status: No .onion sites in database")
            print("   Action: Add sites via:")
            print("      POST /api/darkweb/sites with {group_name, url, ...}")
            print("   Or: POST /api/breach/excel/import to load from OneDrive")
        else:
            print(f"✅ {onion_count} onion sites configured")
            
            # Check scan status
            async with db._conn.execute(
                """SELECT 
                   last_status,
                   COUNT(*) as cnt
                FROM onion_sites
                GROUP BY last_status"""
            ) as cur:
                statuses = await cur.fetchall()
            
            status_map = {row["last_status"]: row["cnt"] for row in statuses}
            online = status_map.get("online", 0)
            offline = status_map.get("offline", 0)
            pending = sum(status_map.get(k, 0) for k in [None, "pending"])
            
            print(f"   • Online:   {online}")
            print(f"   • Offline:  {offline}")
            print(f"   • Pending:  {pending}")
            
            if online == 0 and offline > 0:
                print("   ⚠️  All sites marked OFFLINE - onion_monitor might not be running with Tor")
                print("   Action:")
                print("      1. Check Tor status: /api/darkweb/tor/status")
                print("      2. If Tor not running, enable: ENABLE_DARKWEB=true")
                print("      3. Manually trigger: POST /api/onion/scan")
        
        # Check last scan time
        async with db._conn.execute(
            "SELECT MAX(last_checked) as last_check FROM onion_sites"
        ) as cur:
            row = await cur.fetchone()
            last_check = row["last_check"] if row else None
        
        if last_check:
            print(f"   Last scan: {last_check}")
    except Exception as e:
        print(f"❌ Error checking onion sites: {e}")
    
    print()
    
    # ===== ISSUE 3: Breach Markets not getting new data =====
    print("🛒  ISSUE 3: BREACH MARKETS - NOT GETTING NEW DATA")
    print("-" * 80)
    try:
        async with db._conn.execute(
            "SELECT COUNT(*) as cnt FROM breach_markets"
        ) as cur:
            row = await cur.fetchone()
            market_count = row["cnt"] if row else 0
        
        if market_count == 0:
            print("❌ NO BREACH MARKETS IN DATABASE!")
            print("   Status: RansomLook sync hasn't run yet")
            print("   Action:")
            print("      1. Wait for ransomlook_market job (runs every 10 min)")
            print("      2. Or manually trigger: POST /api/breach/markets/refresh")
        else:
            print(f"✅ {market_count:,} breach markets tracked")
            
            # Check market status
            async with db._conn.execute(
                """SELECT 
                   COUNT(*) as cnt,
                   SUM(CASE WHEN last_status = '200' THEN 1 ELSE 0 END) as online,
                   SUM(CASE WHEN last_status != '200' AND last_status IS NOT NULL THEN 1 ELSE 0 END) as offline
                FROM breach_markets"""
            ) as cur:
                row = await cur.fetchone()
            
            total = row["cnt"] or 0
            online = row["online"] or 0
            offline = row["offline"] or 0
            pending = total - online - offline
            
            print(f"   • Total:   {total:,}")
            print(f"   • Online:  {online} ({100*online/total:.1f}%)")
            print(f"   • Offline: {offline} ({100*offline/total:.1f}%)")
            print(f"   • Pending: {pending}")
            
            # Check last refresh time
            async with db._conn.execute(
                "SELECT MAX(updated_at) as last_update FROM breach_markets"
            ) as cur:
                row = await cur.fetchone()
                last_update = row["last_update"] if row else None
            
            if last_update:
                print(f"   Last update: {last_update}")
            
            if offline > total * 0.8:
                print("   ⚠️  >80% markets marked OFFLINE - check connectivity")
                print("   Action:")
                print("      1. Check network connectivity")
                print("      2. Verify RansomlookMarketConnector is fetching URLs")
                print("      3. Manually refresh: POST /api/breach/markets/refresh")
    except Exception as e:
        print(f"❌ Error checking breach markets: {e}")
    
    print()
    
    # ===== ISSUE 4: Excel export missing dashboards =====
    print("📊 ISSUE 4: EXCEL EXPORT - MISSING DARK WEB & BREACH MARKET SHEETS")
    print("-" * 80)
    from pathlib import Path
    import os
    
    export_path = Path(getattr(settings, "EXCEL_EXPORT_PATH", "data/exports"))
    export_file = export_path / "threatintel_export.xlsx"
    
    if export_file.exists():
        file_size = os.path.getsize(export_file)
        mod_time = os.path.getmtime(export_file)
        print(f"✅ Excel file exists: {export_file}")
        print(f"   Size: {file_size:,} bytes")
        
        # Try to read sheet names
        try:
            from openpyxl import load_workbook
            wb = load_workbook(export_file)
            sheets = wb.sheetnames
            print(f"   Sheets: {', '.join(sheets)}")
            
            has_darkweb = any('onion' in s.lower() or 'dark' in s.lower() for s in sheets)
            has_breach = any('breach' in s.lower() or 'market' in s.lower() for s in sheets)
            
            if not has_darkweb:
                print("   ❌ Missing: Dark Web Manager sheet")
            if not has_breach:
                print("   ❌ Missing: Breach Markets sheet")
        except Exception as e:
            print(f"   ⚠️  Could not read sheets: {e}")
    else:
        print(f"❌ Excel file NOT FOUND: {export_file}")
        print(f"   Status: First export hasn't run yet")
        print(f"   Action: Wait for excel_export job (runs every 60 min now)")
        print(f"   Or trigger manually via: GET /api/export/excel")
    
    print()
    print("=" * 80)
    print("RECOMMENDATIONS")
    print("=" * 80)
    print("""
1. ICS Advisories (missing fields):
   - Wait 60+ minutes for ics_sync job to run
   - Or manually: POST /api/advisory/ics/sync
   - Check if RAPIDAPI_KEY is set in .env
   
2. Dark Web Manager (not scanning):
   - Ensure ENABLE_DARKWEB=true in .env
   - Check Tor is running: /api/darkweb/tor/status
   - Trigger scan: POST /api/onion/scan
   
3. Breach Markets (no new data):
   - Check scheduler is running
   - Manually trigger: POST /api/breach/markets/refresh
   - Verify internet connectivity for ransomlook.io
   
4. Excel Export (missing sheets):
   - Update excel_exporter.py to add:
      - _write_darkweb_sheet() for onion sites
      - _write_breach_markets_sheet() for breach markets
   - Regenerate: GET /api/export/excel
""")
    
    await db.close()

asyncio.run(diagnose())
