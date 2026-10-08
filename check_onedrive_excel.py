import asyncio
from connectors.excel_ingestion import _download_excel_bytes
import openpyxl
import io

async def main():
    data = await _download_excel_bytes()
    if not data:
        print("No data downloaded")
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "ICS-CERT_ADV" in wb.sheetnames:
        sheet = wb["ICS-CERT_ADV"]
    elif len(wb.sheetnames) >= 3:
        sheet = wb.worksheets[2]
    else:
        sheet = wb.worksheets[0]
        
    for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        print("Headers:", row)

asyncio.run(main())
