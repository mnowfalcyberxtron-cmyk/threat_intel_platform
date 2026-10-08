"""
connectors/excel_ingestion.py - OneDrive Excel ingestion for ThreatIntel TIP.

The shared workbook is treated as the live source for three sections:
  Sheet 1: threat actor/group .onion sites -> onion_sites
  Sheet 2: breach markets                  -> breach_markets
  Sheet 3: ICS CVE advisories              -> /api/advisory/ics data merge
"""

import io
import logging
import re
import ssl
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import aiohttp
from config import settings

logger = logging.getLogger("connectors.excel_ingestion")

ONEDRIVE_SHARE_URL = getattr(
    settings,
    "EXCEL_INGESTION_ONEDRIVE_URL",
    "https://1drv.ms/x/c/58897f49075b1bc5/IQCbQQAZ3G0DQ7m06TMh0ECXAeugC9AdA_q_3AWMmc2gyVc?e=Zf9Hoh&nav=MTVfezAwMDAwMDAwLTAwMDEtMDAwMC0wMDAwLTAwMDAwMDAwMDAwMH0",
)
LOCAL_EXCEL_PATH = getattr(settings, "EXCEL_INGESTION_LOCAL_PATH", "")

_NAME_COLS = {
    "name", "group", "group_name", "site_name", "market_name",
    "threat_actor", "actor", "threat_actor_name", "ransomware_group",
    "title",
}
_URL_COLS = {
    "url", "link", "onion_url", "site_url", "address", "domain",
    "market_url", "breach_market_url", "tor_url",
}
_DESC_COLS = {"description", "desc", "notes", "note", "info", "remarks"}
_TYPE_COLS = {"type", "site_type", "category", "kind"}
_STATUS_COLS = {"status", "live_status", "last_status", "active", "availability", "state"}

_ICS_COLS = {
    "cve_id": {"cve", "cve_id", "cveid", "cve_number", "cve_no", "cve_identifier"},
    "advisory_id": {
        "advisory", "advisory_id", "ics_advisory_id", "icsa", "icsa_id",
        "ics_cert_number", "ics_cert_no", "ics_number", "ics_id",
    },
    "published_date": {
        "published", "published_date", "publish_date", "release_date",
        "original_release_date", "date", "advisory_date",
    },
    "last_updated": {"last_updated", "updated", "update_date", "revised", "revision_date"},
    "title": {"title", "advisory_title", "ics_cert_advisory_title", "vulnerability"},
    "patch_availability": {"patch", "patch_available", "patch_availability", "patch_status"},
    "poc_availability": {
        "poc", "poc_available", "poc_availability", "proof_of_concept",
        "exploit_available", "exploit",
    },
    "impact": {"impact", "vulnerability_impact", "effect"},
    "cvss_score": {"cvss", "cvss_score", "score", "base_score", "cumulative_cvss"},
    "cvss_severity": {"severity", "cvss_severity", "cvss_rating", "risk"},
    "affected_vendor": {"vendor", "affected_vendor", "manufacturer"},
    "affected_application": {"product", "application", "affected_application", "affected_product"},
    "affected_version": {
        "affected_version", "affected_versions", "products_affected",
        "affected_products", "vulnerable_version", "vulnerable_versions",
    },
    "fixed_version": {"fixed_version", "fixed_versions", "fixed_in", "patched_version"},
    "cwe": {"cwe", "cwe_number", "cwe_id"},
    "sector": {"sector", "critical_infrastructure_sector", "industry"},
    "advisory_url": {"advisory_url", "url", "link", "source_url", "hyperlink"},
}

_URL_RE = re.compile(r"(https?://[^\s<>'\"]+|[a-z2-7]{16,80}\.onion[^\s<>'\"]*)", re.I)
_CVE_RE = re.compile(r"\bCVE-\d{4}-\d{4,7}\b", re.I)
_ICSA_RE = re.compile(r"\bICSA-\d{2}-\d{3}-\d{2}\b", re.I)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _col_map(headers: list) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for i, header in enumerate(headers):
        norm = _norm_header(header)
        if norm in _NAME_COLS and "name" not in mapping:
            mapping["name"] = i
        elif norm in _URL_COLS and "url" not in mapping:
            mapping["url"] = i
        elif norm in _DESC_COLS and "desc" not in mapping:
            mapping["desc"] = i
        elif norm in _TYPE_COLS and "type" not in mapping:
            mapping["type"] = i
        elif norm in _STATUS_COLS and "status" not in mapping:
            mapping["status"] = i
    return mapping


def _ics_col_map(headers: list) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for i, header in enumerate(headers):
        norm = _norm_header(header)
        for role, aliases in _ICS_COLS.items():
            if norm in aliases and role not in mapping:
                mapping[role] = i
                break
    return mapping


def _cell(row, idx: Optional[int], default: str = "") -> str:
    if idx is None or idx >= len(row):
        return default
    value = row[idx]
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() if value is not None else default


def _normalise_url(url: str, onion_only: bool = False) -> str:
    url = (url or "").strip().strip(".,;)")
    if not url:
        return ""
    if ".onion" in url.lower() and not re.match(r"https?://", url, re.I):
        url = "http://" + url
    if not re.match(r"https?://", url, re.I):
        return ""
    if onion_only and ".onion" not in url.lower():
        return ""
    return url.rstrip("/")


def _extract_url(cells: list, onion_only: bool = False) -> str:
    for value in cells:
        text = str(value or "").strip()
        if not text:
            continue
        for match in _URL_RE.findall(text):
            url = _normalise_url(match, onion_only=onion_only)
            if url:
                return url
    return ""


def _status_to_db(value: str) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return "pending"
    if text in {"200", "ok", "up", "live", "online", "active", "reachable", "working", "yes", "true"}:
        return "200"
    if text in {"pending", "unknown", "unchecked", "new", "not checked"}:
        return "pending"
    if any(token in text for token in ("offline", "inactive", "dead", "down", "unreachable", "timeout", "error", "no", "false")):
        return "offline"
    return text[:80]


def _fallback_name(cells: list, url: str, default: str = "") -> str:
    url_l = url.lower()
    for value in cells:
        text = str(value or "").strip()
        if not text:
            continue
        tl = text.lower()
        if tl in _STATUS_COLS or tl in {"url", "link", "description", "notes"}:
            continue
        if ".onion" in tl or "http://" in tl or "https://" in tl or tl in url_l:
            continue
        if _CVE_RE.search(text) or _ICSA_RE.search(text):
            continue
        return text[:120]
    return default or url


def _parse_date_parts(value: Any) -> Tuple[Optional[int], Optional[int]]:
    if isinstance(value, datetime):
        return value.year, value.month
    if isinstance(value, date):
        return value.year, value.month
    text = str(value or "").strip()
    if not text:
        return None, None
    text = text.replace("/", "-").split("T")[0]
    for pattern, y_idx, m_idx in [
        (r"^(\d{4})-(\d{1,2})-(\d{1,2})", 1, 2),
        (r"^(\d{1,2})-(\d{1,2})-(\d{4})", 3, 1),
    ]:
        match = re.match(pattern, text)
        if match:
            return int(match.group(y_idx)), int(match.group(m_idx))
    months = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    match = re.match(r"^([A-Za-z]+)\s+\d{1,2},?\s+(\d{4})", text)
    if match and match.group(1)[:3].lower() in months:
        return int(match.group(2)), months[match.group(1)[:3].lower()]
    match = re.match(r"^\d{1,2}\s+([A-Za-z]+),?\s+(\d{4})", text)
    if match and match.group(1)[:3].lower() in months:
        return int(match.group(2)), months[match.group(1)[:3].lower()]
    return None, None


def _extract_cve_year(cve_id: str) -> Optional[int]:
    match = re.match(r"(?i)CVE-(\d{4})-", str(cve_id or ""))
    return int(match.group(1)) if match else None


async def _download_excel_bytes() -> Optional[bytes]:
    """Prefer local workbook path if configured, otherwise download from OneDrive."""
    if LOCAL_EXCEL_PATH:
        local_path = Path(LOCAL_EXCEL_PATH)
        if local_path.exists() and local_path.is_file():
            try:
                return local_path.read_bytes()
            except Exception as exc:
                logger.error("[ExcelIngestion] Failed to read local Excel path %s: %s", LOCAL_EXCEL_PATH, exc)
                return None
        logger.warning("[ExcelIngestion] Configured local Excel path does not exist: %s", LOCAL_EXCEL_PATH)

    if not ONEDRIVE_SHARE_URL:
        logger.error("[ExcelIngestion] No Excel ingestion source configured.")
        return None

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,*/*;q=0.9",
        "Accept-Language": "en-US,en;q=0.9",
    }
    ssl_ctx = ssl.create_default_context()
    try:
        ssl_ctx.load_verify_locations()
    except Exception:
        pass

    connector = aiohttp.TCPConnector(ssl=ssl_ctx)
    timeout = aiohttp.ClientTimeout(total=60)
    async with aiohttp.ClientSession(connector=connector, timeout=timeout, headers=headers) as sess:
        try:
            async with sess.get(ONEDRIVE_SHARE_URL, allow_redirects=True) as resp:
                if resp.status != 200:
                    logger.error("[ExcelIngestion] OneDrive share page returned HTTP %s", resp.status)
                    return None
                html = await resp.text(errors="replace")
        except Exception as exc:
            logger.error("[ExcelIngestion] Failed to fetch OneDrive share page: %s", exc)
            return None

        download_url = None
        match = re.search(r'"FileGetUrl"\s*:\s*"([^"]+)"', html)
        if match:
            download_url = match.group(1).replace("\\u0026", "&").replace("\\/", "/")
            logger.info("[ExcelIngestion] Extracted FileGetUrl from OneDrive page")

        if not download_url:
            for pattern in [
                r'"downloadUrl"\s*:\s*"([^"]+)"',
                r'href="([^"]+\.xlsx[^"]*)"',
                r"'(https://[^']*download[^']*\.xlsx[^']*)'",
            ]:
                match = re.search(pattern, html, re.IGNORECASE)
                if match:
                    download_url = match.group(1).replace("\\u0026", "&").replace("\\/", "/")
                    break

        if not download_url:
            embed = re.search(r'"embedUrl"\s*:\s*"([^"]+)"', html)
            if embed:
                download_url = re.sub(
                    r"action=embedview",
                    "action=download",
                    embed.group(1).replace("\\u0026", "&").replace("\\/", "/"),
                )

        if not download_url:
            logger.error("[ExcelIngestion] Could not extract a OneDrive download URL.")
            return None

        try:
            async with sess.get(download_url, allow_redirects=True) as dl:
                if dl.status != 200:
                    logger.error("[ExcelIngestion] Download URL returned HTTP %s", dl.status)
                    return None
                content_type = dl.headers.get("Content-Type", "")
                if "html" in content_type.lower():
                    logger.error("[ExcelIngestion] Download returned HTML instead of Excel.")
                    return None
                data = await dl.read()
                logger.info("[ExcelIngestion] Downloaded Excel file: %d bytes", len(data))
                return data
        except Exception as exc:
            logger.error("[ExcelIngestion] Failed to download Excel file: %s", exc)
            return None


def _parse_site_sheet(ws, onion_only: bool, default_type: str) -> List[Dict[str, str]]:
    rows_out: List[Dict[str, str]] = []
    col_map: Dict[str, int] = {}
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = [c for c in row]
        if not header_seen:
            candidate_map = _col_map(cells)
            row_url = _extract_url(cells, onion_only=onion_only)
            if candidate_map.get("url") is not None and not row_url:
                col_map = candidate_map
                header_seen = True
                continue
            header_seen = True

        url = _normalise_url(_cell(cells, col_map.get("url")), onion_only=onion_only)
        if not url:
            url = _extract_url(cells, onion_only=onion_only)
        if not url:
            continue

        rows_out.append({
            "name": _cell(cells, col_map.get("name")) or _fallback_name(cells, url),
            "url": url,
            "description": _cell(cells, col_map.get("desc")),
            "site_type": _cell(cells, col_map.get("type")) or default_type,
            "last_status": _status_to_db(_cell(cells, col_map.get("status"))),
        })
    return rows_out


def _parse_ics_sheet(ws) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []
    col_map: Dict[str, int] = {}
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = [c for c in row]
        if not header_seen:
            candidate_map = _ics_col_map(cells)
            has_cve = any(_CVE_RE.search(str(c or "")) for c in cells)
            has_advisory = any(_ICSA_RE.search(str(c or "")) for c in cells)
            if candidate_map and not (has_cve or has_advisory):
                col_map = candidate_map
                header_seen = True
                continue
            header_seen = True

        raw_cves = _cell(cells, col_map.get("cve_id"))
        cves = _CVE_RE.findall(raw_cves)
        if not cves:
            for cell in cells:
                cves.extend(_CVE_RE.findall(str(cell or "")))
        cves = list(dict.fromkeys(cve.upper() for cve in cves))
        if not cves:
            continue

        advisory_id = _cell(cells, col_map.get("advisory_id"))
        if not advisory_id:
            for cell in cells:
                match = _ICSA_RE.search(str(cell or ""))
                if match:
                    advisory_id = match.group(0).upper()
                    break

        published_raw = cells[col_map["published_date"]] if "published_date" in col_map and col_map["published_date"] < len(cells) else ""
        updated_raw = cells[col_map["last_updated"]] if "last_updated" in col_map and col_map["last_updated"] < len(cells) else ""
        published = _cell(cells, col_map.get("published_date"))
        updated = _cell(cells, col_map.get("last_updated"))
        pub_year, pub_month = _parse_date_parts(published_raw or published)
        upd_year, upd_month = _parse_date_parts(updated_raw or updated)

        advisory_url = _normalise_url(_cell(cells, col_map.get("advisory_url")))
        if not advisory_url and advisory_id:
            advisory_url = f"https://www.cisa.gov/news-events/ics-advisories/{advisory_id.lower()}"

        for cve in cves:
            cve_year = _extract_cve_year(cve)
            rows_out.append({
                "ics_number": advisory_id,
                "advisory_id": advisory_id,
                "cve_id": cve,
                "title": _cell(cells, col_map.get("title")),
                "patch_availability": _cell(cells, col_map.get("patch_availability")),
                "poc_availability": _cell(cells, col_map.get("poc_availability")),
                "impact": _cell(cells, col_map.get("impact")),
                "cvss_score": _cell(cells, col_map.get("cvss_score")),
                "cvss_severity": _cell(cells, col_map.get("cvss_severity")),
                "affected_vendor": _cell(cells, col_map.get("affected_vendor")),
                "affected_application": _cell(cells, col_map.get("affected_application")),
                "affected_version": _cell(cells, col_map.get("affected_version")),
                "fixed_version": _cell(cells, col_map.get("fixed_version")),
                "cwe": _cell(cells, col_map.get("cwe")),
                "sector": _cell(cells, col_map.get("sector")),
                "release_date": published,
                "published_date": published,
                "last_updated": updated,
                "published_year": pub_year,
                "published_month": pub_month,
                "year": cve_year or pub_year or upd_year,
                "month": pub_month or upd_month,
                "cve_year": cve_year,
                "release_year": pub_year,
                "release_month": pub_month,
                "update_year": upd_year,
                "update_month": upd_month,
                "advisory_url": advisory_url,
                "data_source": "excel",
            })
    return rows_out


def _parse_excel(data: bytes) -> Tuple[list, list, list]:
    """Parse workbook bytes into sheet1 onion rows, sheet2 market rows, sheet3 ICS CVE rows."""
    try:
        import openpyxl
    except ImportError:
        logger.error("[ExcelIngestion] openpyxl is not installed.")
        return [], [], []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        logger.error("[ExcelIngestion] Failed to parse Excel workbook: %s", exc)
        return [], [], []

    sheet1_data: list = []
    sheet2_data: list = []
    sheet3_data: list = []
    sheet_names = wb.sheetnames
    logger.info("[ExcelIngestion] Sheets found: %s", sheet_names)

    # Heuristic mapping
    onion_sheet_name = None
    market_sheet_name = None
    ics_sheet_name = None

    for name in sheet_names:
        sh = wb[name]
        headers = []
        for row in sh.iter_rows(max_row=5, values_only=True):
            if row and not all(c is None or str(c).strip() == "" for c in row):
                headers = [str(c).strip() for c in row if c is not None]
                break

        name_l = name.lower()
        is_ics = "ics" in name_l or "cve" in name_l or "advisory" in name_l or "cisa" in name_l
        if not is_ics:
            ics_map = _ics_col_map(headers)
            if "cve_id" in ics_map or "advisory_id" in ics_map:
                is_ics = True

        if is_ics:
            ics_sheet_name = name
            continue

        is_onion = "onion" in name_l or "tor" in name_l or "leak" in name_l or "group" in name_l
        if is_onion:
            onion_sheet_name = name
            continue

        is_market = "market" in name_l or "forum" in name_l or "breach" in name_l
        if is_market:
            market_sheet_name = name
            continue

    # Fallback by index for unassigned slots
    used_names = {onion_sheet_name, market_sheet_name, ics_sheet_name}
    remaining_names = [n for n in sheet_names if n not in used_names]

    if not onion_sheet_name:
        if remaining_names:
            onion_sheet_name = remaining_names.pop(0)
        elif len(sheet_names) >= 1:
            onion_sheet_name = sheet_names[0]

    if not market_sheet_name:
        if remaining_names:
            market_sheet_name = remaining_names.pop(0)
        elif len(sheet_names) >= 2:
            market_sheet_name = sheet_names[1]

    if not ics_sheet_name:
        if remaining_names:
            ics_sheet_name = remaining_names.pop(0)
        elif len(sheet_names) >= 3:
            ics_sheet_name = sheet_names[2]

    logger.info("[ExcelIngestion] Mapped sheets: S1(Onion)=%s, S2(Market)=%s, S3(ICS)=%s",
                onion_sheet_name, market_sheet_name, ics_sheet_name)

    if onion_sheet_name:
        sheet1_data = _parse_site_sheet(wb[onion_sheet_name], onion_only=True, default_type="ransomware")
        logger.info("[ExcelIngestion] Sheet1 onion rows (%s): %d", onion_sheet_name, len(sheet1_data))
    if market_sheet_name:
        sheet2_data = _parse_site_sheet(wb[market_sheet_name], onion_only=False, default_type="market")
        logger.info("[ExcelIngestion] Sheet2 breach market rows (%s): %d", market_sheet_name, len(sheet2_data))
    if ics_sheet_name:
        sheet3_data = _parse_ics_sheet(wb[ics_sheet_name])
        logger.info("[ExcelIngestion] Sheet3 ICS CVE rows (%s): %d", ics_sheet_name, len(sheet3_data))

    wb.close()
    return sheet1_data, sheet2_data, sheet3_data


async def _upsert_onion_sites(db, rows: list) -> int:
    added = 0
    updated = 0
    for row in rows:
        url = row["url"]
        name = row["name"] or url
        desc = row["description"] or ""
        status = row.get("last_status") or "pending"
        site_type = row.get("site_type") or "ransomware"
        now = _now()
        try:
            async with db._conn.execute("SELECT id FROM onion_sites WHERE url=?", (url,)) as cur:
                existing = await cur.fetchone()
            if existing:
                if status != "pending":
                    await db._conn.execute(
                        """UPDATE onion_sites
                           SET group_name=?, description=CASE WHEN ? != '' THEN ? ELSE description END,
                               site_type=?, active=1, last_status=?, last_checked=?
                           WHERE url=?""",
                        (name, desc, desc, site_type, status, now, url),
                    )
                else:
                    await db._conn.execute(
                        """UPDATE onion_sites
                           SET group_name=?, description=CASE WHEN ? != '' THEN ? ELSE description END,
                               site_type=?, active=1
                           WHERE url=?""",
                        (name, desc, desc, site_type, url),
                    )
                updated += 1
            else:
                await db._conn.execute(
                    """INSERT INTO onion_sites
                       (group_name, url, description, active, site_type, created_at, last_status)
                       VALUES (?, ?, ?, 1, ?, ?, ?)""",
                    (name, url, desc, site_type, now, status),
                )
                added += 1
                logger.info("[ExcelIngestion] Added onion site: %s", url)
        except Exception as exc:
            logger.warning("[ExcelIngestion] Error upserting onion site %s: %s", url, exc)
    await db._conn.commit()
    logger.info("[ExcelIngestion] Onion sheet upsert: +%d added, %d updated", added, updated)
    return added


async def _upsert_breach_markets(db, rows: list) -> int:
    added = 0
    updated = 0
    for row in rows:
        url = row["url"]
        name = row["name"] or url
        desc = row["description"] or ""
        status = row.get("last_status") or "pending"
        site_type = row.get("site_type") or "market"
        now = _now()
        try:
            async with db._conn.execute("SELECT id FROM breach_markets WHERE url=?", (url,)) as cur:
                existing = await cur.fetchone()
            if existing:
                if status != "pending":
                    await db._conn.execute(
                        """UPDATE breach_markets
                           SET name=?, description=CASE WHEN ? != '' THEN ? ELSE description END,
                               site_type=?, source='excel_import', active=1,
                               last_status=?, last_checked=?, updated_at=?
                           WHERE url=?""",
                        (name, desc, desc, site_type, status, now, now, url),
                    )
                else:
                    await db._conn.execute(
                        """UPDATE breach_markets
                           SET name=?, description=CASE WHEN ? != '' THEN ? ELSE description END,
                               site_type=?, source='excel_import', active=1, updated_at=?
                           WHERE url=?""",
                        (name, desc, desc, site_type, now, url),
                    )
                updated += 1
            else:
                await db._conn.execute(
                    """INSERT INTO breach_markets
                       (name, url, description, site_type, source, active, last_status, created_at, updated_at)
                       VALUES (?, ?, ?, ?, 'excel_import', 1, ?, ?, ?)""",
                    (name, url, desc, site_type, status, now, now),
                )
                added += 1
                logger.info("[ExcelIngestion] Added breach market: %s", url)
        except Exception as exc:
            logger.warning("[ExcelIngestion] Error upserting breach market %s: %s", url, exc)
    await db._conn.commit()
    logger.info("[ExcelIngestion] Breach market sheet upsert: +%d added, %d updated", added, updated)
    return added


async def _upsert_ics_advisories(db, rows: list) -> int:
    added = 0
    for row in rows:
        try:
            # Map excel sheet3 keys to database schema fields
            record = {
                "ics_number": row.get("ics_number") or row.get("advisory_id") or "",
                "cve_id": row.get("cve_id") or "N/A",
                "title": row.get("title") or "",
                "affected_vendor": row.get("affected_vendor") or "",
                "affected_application": row.get("affected_application") or "",
                "cvss_score": row.get("cvss_score") or "",
                "cvss_severity": row.get("cvss_severity") or "medium",
                "release_date": row.get("release_date") or row.get("published_date") or "",
                "last_updated": row.get("last_updated") or "",
                "advisory_url": row.get("advisory_url") or row.get("url") or "",
                "sector": row.get("sector") or "",
                "patch_availability": row.get("patch_availability") or "Unknown",
                "poc_availability": row.get("poc_availability") or "Unknown",
                "impact": row.get("impact") or "",
                "affected_version": row.get("affected_version") or "",
                "fixed_version": row.get("fixed_version") or "",
                "cwe": row.get("cwe") or "",
                "data_source": "excel",
            }
            _, is_new = await db.upsert_ics_advisory(record)
            if is_new:
                added += 1
        except Exception as exc:
            logger.warning("[ExcelIngestion] Error upserting ICS row from Excel: %s", exc)
    logger.info("[ExcelIngestion] ICS sheet upsert: +%d added", added)
    return added


async def fetch_excel_ics_rows() -> List[Dict[str, Any]]:
    """Download the workbook and return normalized Sheet 3 ICS CVE rows."""
    data = await _download_excel_bytes()
    if not data:
        return []
    _, _, sheet3_rows = _parse_excel(data)
    return sheet3_rows


async def run_excel_ingestion(db) -> Dict[str, Any]:
    """Download workbook, upsert sheet 1/2/3 into DB."""
    logger.info("[ExcelIngestion] Starting OneDrive Excel ingestion (3-sheet mode)...")
    result = {
        "status": "ok",
        "onion_added": 0,
        "markets_added": 0,
        "ics_added": 0,
        "onion_total_rows": 0,
        "markets_total_rows": 0,
        "ics_total_rows": 0,
        "error": None,
    }

    data = await _download_excel_bytes()
    if not data:
        result["status"] = "error"
        result["error"] = "Failed to download Excel file from OneDrive"
        logger.warning("[ExcelIngestion] Skipping - could not download file.")
        return result

    sheet1_rows, sheet2_rows, sheet3_rows = _parse_excel(data)
    result["onion_total_rows"] = len(sheet1_rows)
    result["markets_total_rows"] = len(sheet2_rows)
    result["ics_total_rows"] = len(sheet3_rows)

    if sheet1_rows:
        result["onion_added"] = await _upsert_onion_sites(db, sheet1_rows)
    if sheet2_rows:
        result["markets_added"] = await _upsert_breach_markets(db, sheet2_rows)
    if sheet3_rows:
        result["ics_added"] = await _upsert_ics_advisories(db, sheet3_rows)

    logger.info(
        "[ExcelIngestion] Done. Onion S1: +%d/%d. Markets S2: +%d/%d. ICS S3: +%d/%d rows.",
        result["onion_added"], result["onion_total_rows"],
        result["markets_added"], result["markets_total_rows"],
        result["ics_added"], result["ics_total_rows"],
    )
    return result
